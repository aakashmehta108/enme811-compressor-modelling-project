"""Build the Word technical report from verified CSV/JSON outputs."""
from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
FIG = ROOT / "report" / "figures"
OUT = ROOT / "report" / "technical_report.docx"
values = json.loads((RESULTS / "report_values.json").read_text())
candidates = pd.read_csv(RESULTS / "candidate_model_comparison.csv")
one_d = pd.read_csv(RESULTS / "one_dimensional_fits.csv")
new_cases = pd.read_csv(RESULTS / "new_case_predictions.csv")
calc = pd.read_csv(RESULTS / "calculated_variables.csv")

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.75)
section.right_margin = Inches(0.75)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.07
for name, size, color in [
    ("Heading 1", 15, "17365D"),
    ("Heading 2", 12.5, "1F4E79"),
    ("Heading 3", 11, "2E5A88"),
]:
    st = styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 9)
    st.paragraph_format.space_after = Pt(5)

# Footer with page number field.
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("ENME811 Compressor Modelling Project  |  Page ")
fld_char1 = OxmlElement("w:fldChar"); fld_char1.set(qn("w:fldCharType"), "begin")
instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = "PAGE"
fld_char2 = OxmlElement("w:fldChar"); fld_char2.set(qn("w:fldCharType"), "end")
r = OxmlElement("w:r"); r.append(fld_char1); fp._p.append(r)
r = OxmlElement("w:r"); r.append(instr); fp._p.append(r)
r = OxmlElement("w:r"); r.append(fld_char2); fp._p.append(r)


def p(text="", style=None, bold=False, italic=False, align=None, size=None):
    par = doc.add_paragraph(style=style)
    if align is not None:
        par.alignment = align
    if text:
        r = par.add_run(text)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    return par


def add_table(headers, rows, widths=None, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = str(h)
        for par in hdr[i].paragraphs:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in par.runs:
                r.bold = True
                r.font.size = Pt(font_size)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "D9EAF3")
        hdr[i]._tc.get_or_add_tcPr().append(shading)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
            for par in cells[i].paragraphs:
                for r in par.runs:
                    r.font.size = Pt(font_size)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()
    return table


def add_fig(filename, caption, width=6.2):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(str(FIG / filename), width=Inches(width))
    cap = p(caption, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9)
    return cap


def equation(text):
    par = p(text, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
    par.paragraph_format.space_before = Pt(4)
    par.paragraph_format.space_after = Pt(8)
    return par

# Cover
p("Auckland University of Technology", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=16)
p("Department of Mechanical Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
p("ENME811 Advanced CAD/CAM Applications", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14)
p("Compressor Modelling Project", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=22)
p("MATLAB Modelling of a Variable-Speed Hermetic Reciprocating Compressor", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
add_table(["Item", "Detail"], [
    ["Refrigerant", "R134a"],
    ["Student name", "______________________________"],
    ["Student ID", "______________________________"],
    ["Submission date", "October 2026"],
    ["Deliverables", "MATLAB package, technical report and CSV data"],
], widths=[1.7, 4.8])
p("This report presents the compressor-selection basis, thermodynamic calculations, polynomial regression models, MATLAB implementation, independent validation and engineering discussion required by the ENME811 project brief.", align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

p("Executive summary", style="Heading 1")
p("A MATLAB model was developed for a variable-speed hermetic reciprocating compressor operating with R134a. The workflow covers operating-point preparation, CoolProp refrigerant-property calculation, volumetric and isentropic efficiency evaluation, fixed-frequency and two-dimensional polynomial regression, performance prediction and independent validation.")
p(f"The client-supplied design duty is 50 kW cooling with R134a at an evaporating temperature of −10 °C and 5 K suction superheat. The selected reference compressor has a displacement of {values['Vd_cm3_rev']:.1f} cm³/rev, sized so the fitted model predicts 50.02 kW at the design point (45 °C condensing, 50 Hz), and is evaluated over evaporating temperatures from −15 to 10 °C, condensing temperatures from 35 to 55 °C and frequencies from 30 to 75 Hz, with 5 K suction superheat. The dataset contains 84 points: 54 fitting points and 30 independent validation points. Pressure ratio ranges from 2.14 to 9.10, volumetric efficiency from 0.462 to 0.865, and isentropic efficiency from 0.530 to 0.797.")
equation("ηv = 0.888934 − 0.039473β − 0.001200β² + 0.093179fn − 0.001939βfn − 0.033499fn²")
equation("ηis = 0.752633 + 0.009202β − 0.003705β² + 0.040816fn − 0.003010βfn − 0.009658fn²")
p("On the fitting data, the volumetric model gives R² = 0.9974 and RMSE = 0.00506; the isentropic model gives R² = 0.9837 and RMSE = 0.00798. On the 30 independent validation points, mass-flow MAPE is 0.60%, power MAPE is 0.76%, and discharge-temperature MAE is 0.88 °C. These low errors reflect the smooth representative dataset and should not be interpreted as guaranteed accuracy for a physical compressor without measured manufacturer or laboratory data.")

p("1. Introduction, objectives and scope", style="Heading 1")
p("The compressor determines refrigerant mass flow, pressure lift and most of the electrical power demand in a vapour-compression system. For a variable-speed hermetic reciprocating compressor, a fixed-speed curve is insufficient because speed changes both displacement rate and loss mechanisms. A compact polynomial efficiency model is useful for design studies, control development and system-level simulation.")
p("The objectives are to select and justify an operating envelope; organise fitting and validation data; calculate refrigerant states and efficiencies using CoolProp; fit one-dimensional and two-dimensional polynomial correlations; predict mass flow, power input and discharge temperature; and quantify independent validation errors. The scope is steady-state compressor behaviour. Suction/discharge line pressure drops, transient inverter behaviour, oil circulation and heat-exchanger dynamics are outside the compressor sub-model.")

p("2. Operating conditions, compressor selection and data sources", style="Heading 1")
p("2.1 Adopted operating conditions", style="Heading 2")
p("The operating conditions for this study were supplied for the commission: refrigerant R134a, a design cooling capacity of 50 kW, evaporating temperature −10 °C and suction superheat 5 K. The condensing temperature was not specified, so a design condensing temperature of 45 °C was adopted, and the modelling envelope spans 35–55 °C. The selected displacement (1,147.6 cm³/rev) is sized so the fitted model predicts 50.02 kW at the design point (45 °C condensing, 50 Hz).")
add_table(["Variable", "Adopted range or value"], [
    ["Refrigerant", "R134a (supplied)"],
    ["Design cooling capacity", "50 kW (supplied)"],
    ["Design evaporating temperature", "−10 °C (supplied)"],
    ["Design condensing temperature", "45 °C (adopted); envelope 35, 45, 55 °C"],
    ["Evaporating-temperature envelope", "−15, −10, −5, 0, 5, 10 °C"],
    ["Fitting frequencies", "30, 45, 60 Hz"],
    ["Validation frequencies", "75 Hz plus off-grid 38 and 68 Hz cases"],
    ["Suction superheat", "5 K (supplied)"],
    ["Suction pressure", "1.639–4.146 bar"],
    ["Discharge pressure", "8.870–14.915 bar"],
    ["Pressure ratio", "2.139–9.098"],
], widths=[2.0, 4.5])

p("2.2 Selected reference compressor", style="Heading 2")
add_table(["Parameter", "Selected value"], [
    ["Type", "Hermetic reciprocating, variable speed"],
    ["Reference designation", "VS-HR-1150"],
    ["Refrigerant compatibility", "R134a"],
    ["Bore × stroke × cylinders", "70 mm × 49.7 mm × 6"],
    ["Displacement", "1,147.6 cm³/rev"],
    ["Speed range represented", "1,800–4,500 rpm (30–75 Hz)"],
], widths=[2.1, 4.4])
p("The displacement was sized from the supplied duty. At the design point (Tevap = −10 °C, Tcond = 45 °C, 5 K superheat, 50 Hz) the refrigerating effect is 132.98 kJ/kg, so 50 kW requires 0.376 kg/s; with the modelled volumetric efficiency at β = 5.78 the required displacement is approximately 1,148 cm³/rev, realised as a six-cylinder 70 mm × 49.7 mm geometry. R134a is the supplied refrigerant and has a well-validated CoolProp equation of state. The envelope covers the duty while keeping pressure ratio below approximately 9.1.")

p("2.3 Data source and fitting/validation split", style="Heading 2")
p("Because no manufacturer selection-software export was supplied, the performance table is a documented representative manufacturer-style dataset. CoolProp generated the refrigerant properties, and performance values were produced from physically realistic efficiency maps with small scatter. The generation script is included for auditability. The numerical results therefore describe the adopted reference case; replacing the CSV with a genuine manufacturer export would model a named product without changing the MATLAB code.")
p("The dataset contains 54 fitting points (six evaporating temperatures × three condensing temperatures × three frequencies) and 30 validation points (the same 18 temperature combinations at 75 Hz plus 12 off-grid cases). Validation rows are excluded from coefficient estimation.")

p("3. Literature review and modelling assumptions", style="Heading 1")
p("Polynomial compressor maps are widely used in refrigeration simulation because they are compact, fast and easy to embed in system models. The assignment's key reference, Wang and Lu (2025), follows a streamlined regression approach for speed-driven hermetic reciprocating compressors. The present implementation fits efficiencies rather than mass flow and power directly, retaining CoolProp state calculations and making extrapolation behaviour easier to inspect.")
p("Assumptions are steady-state operation; suction and discharge pressures equal to saturation pressures at the evaporating and condensing temperatures; 5 K suction superheat; negligible line pressure drops; speed N = 60f rpm; and total hermetic electrical input included in the isentropic-efficiency definition. Predictions are intended for the fitted pressure-ratio and frequency envelope.")

p("4. Thermodynamic property and efficiency calculations", style="Heading 1")
p("For each point, CoolProp provides Psuc = Psat(Tevap), Pdis = Psat(Tcond), suction density, suction enthalpy h1, suction entropy s1 and isentropic discharge enthalpy h2s = h(Pdis,s1).")
equation("β = Pdis / Psuc")
equation("ηv = ṁ·vsuc / (Vd·f)")
equation("ηis = ṁ(h2s − h1) / Wel")
equation("ṁpred = ηv,pred·Vd·f / vsuc;    Wpred = ṁpred(h2s − h1)/ηis,pred")
equation("h2,pred = h1 + (h2s − h1)/ηis,pred;    Tdis,pred = T(Pdis,h2,pred)")
p("Across all 84 points, suction density ranges from 8.09 to 19.69 kg/m³, h1 from 393.80 to 409.01 kJ/kg and h2s from 425.22 to 441.23 kJ/kg. Reference mass flow ranges from 0.129 to 1.454 kg/s, power from 9.73 to 48.13 kW and discharge temperature from 46.6 to 104.3 °C.")

p("5. Data preparation and operating-point selection", style="Heading 1")
p("The MATLAB loader checks required columns and missing values. The calculated-variable CSV adds pressure ratio, suction density/specific volume, enthalpies, calculated efficiencies, predicted efficiencies and outputs, and errors. Units are stated in the code and output headers.")
add_table(["Field group", "Contents"], [
    ["Identifiers", "point_id; dataset (Fitting/Validation)"],
    ["Operating conditions", "T_evap_C, T_cond_C, freq_Hz, speed_rpm"],
    ["Reference performance", "m_dot_kg_s, W_el_kW, T_dis_C"],
    ["Calculated variables", "beta, rho_suc, v_suc, h1, s1, h2is, eta_v, eta_is, predictions and errors"],
], widths=[1.8, 4.7])

p("6. Polynomial regression results", style="Heading 1")
p("6.1 One-dimensional fixed-frequency fits", style="Heading 2")
rows = []
for _, r in one_d.iterrows():
    rows.append([
        f"{r.freq_Hz:g} Hz",
        f"{r.eta_v_c0:.6f}, {r.eta_v_c1_beta:.6f}, {r.eta_v_c2_beta2:.6f}",
        f"{r.eta_v_R2:.4f}", f"{r.eta_v_RMSE:.5f}",
        f"{r.eta_is_c0:.6f}, {r.eta_is_c1_beta:.6f}, {r.eta_is_c2_beta2:.6f}",
        f"{r.eta_is_R2:.4f}", f"{r.eta_is_RMSE:.5f}",
    ])
add_table(["Frequency", "ηv coefficients", "R²", "RMSE", "ηis coefficients", "R²", "RMSE"], rows, widths=[0.75,1.45,0.55,0.65,1.45,0.55,0.65], font_size=7.5)

p("6.2 Candidate two-dimensional models", style="Heading 2")
rows = []
labels = {"A": "β only", "B": "β + frequency", "C": "β + frequency + interaction"}
for _, r in candidates.iterrows():
    rows.append([r.model + " — " + labels[r.model], r.terms, f"{r.eta_v_R2:.4f}", f"{r.eta_v_RMSE:.5f}", f"{r.eta_is_R2:.4f}", f"{r.eta_is_RMSE:.5f}"])
add_table(["Model", "Terms", "ηv R²", "ηv RMSE", "ηis R²", "ηis RMSE"], rows, widths=[1.25,2.15,0.75,0.85,0.75,0.85], font_size=8)
p("Adding frequency reduces volumetric-efficiency RMSE from 0.00789 to approximately 0.0051. Model C gives the lowest RMSE for both efficiencies and is retained for prediction.")

p("6.3 Final correlations", style="Heading 2")
equation("ηv = 0.888934 − 0.039473β − 0.001200β² + 0.093179fn − 0.001939βfn − 0.033499fn²")
equation("ηis = 0.752633 + 0.009202β − 0.003705β² + 0.040816fn − 0.003010βfn − 0.009658fn²")
add_table(["Response", "R²", "Adjusted R²", "RMSE"], [
    ["ηv", "0.9974", "0.9971", "0.00506"],
    ["ηis", "0.9837", "0.9820", "0.00798"],
], widths=[1.3, 1.3, 1.5, 1.3])

p("7. MATLAB model structure and implementation", style="Heading 1")
add_table(["File", "Role"], [
    ["main_compressor_model.m", "Runs the complete workflow and writes result tables"],
    ["loadCompressorData.m", "Imports the CSV and checks data quality"],
    ["getRefrigerantProp.m", "Calls CoolProp through MATLAB's Python interface"],
    ["calcEfficiencies.m", "Calculates properties, pressure ratio and efficiencies"],
    ["fitEfficiencyModels.m", "Fits 1-D and 2-D models using fitting rows only"],
    ["predictPerformance.m", "Predicts mass flow, power, discharge temperature and efficiencies"],
    ["validateModel.m", "Independent validation, metrics and figures"],
], widths=[2.1, 4.4])

p("8. Simulation results, validation and error analysis", style="Heading 1")
p("8.1 Efficiency behaviour", style="Heading 2")
add_fig("eta_v_vs_beta.png", "Figure 1. Volumetric efficiency versus pressure ratio.")
p("Volumetric efficiency decreases strongly with pressure ratio, from approximately 0.86 at β ≈ 2.14 to 0.46 at β ≈ 9.10, consistent with clearance-volume re-expansion and leakage effects.")
add_fig("eta_is_vs_beta.png", "Figure 2. Isentropic efficiency versus pressure ratio.")
p("Isentropic efficiency decreases from approximately 0.80 to 0.53. Frequency has a visible but secondary effect compared with pressure ratio.")
add_fig("eta_is_surface.png", "Figure 3. Two-dimensional isentropic-efficiency surface.")

p("8.2 Design-point performance (supplied duty)", style="Heading 2")
add_table(["Quantity", "Value"], [
    ["Evaporating temperature", "−10 °C"],
    ["Design condensing temperature", "45 °C"],
    ["Suction superheat", "5 K"],
    ["Rated frequency", "50 Hz"],
    ["Suction / discharge pressure", "2.006 / 11.599 bar"],
    ["Pressure ratio β", "5.782"],
    ["Refrigerating effect h1 − hf", "132.98 kJ/kg"],
    ["Required mass flow for 50 kW", "0.3760 kg/s"],
    ["Predicted ηv / ηis", "0.669 / 0.696"],
    ["Predicted mass flow", "0.3762 kg/s"],
    ["Predicted electrical power", "20.33 kW"],
    ["Predicted discharge temperature", "71.4 °C"],
    ["Predicted cooling capacity", "50.02 kW"],
    ["Cooling COP", "2.46"],
], widths=[2.4, 3.4])
p("The selected displacement therefore meets the 50 kW duty at the design point with a predicted input of approximately 20.3 kW (cooling COP 2.46).")

p("8.3 Predictions at representative operating points", style="Heading 2")
rows = []
for _, r in new_cases.iterrows():
    rows.append([f"{r.T_evap_C:g}", f"{r.T_cond_C:g}", f"{r.freq_Hz:g}", f"{r.beta:.3f}", f"{r.eta_v_pred:.3f}", f"{r.eta_is_pred:.3f}", f"{r.m_dot_pred_kg_s*1000:.2f}", f"{r.W_pred_kW:.3f}", f"{r.T_dis_pred_C:.1f}"])
add_table(["Tevap °C", "Tcond °C", "f Hz", "β", "ηv", "ηis", "ṁ g/s", "W kW", "Tdis °C"], rows, widths=[0.7,0.7,0.55,0.65,0.6,0.6,0.75,0.65,0.75], font_size=8)

p("8.4 Independent validation", style="Heading 2")
add_table(["Quantity", "MAPE / MAE", "RMSE", "Maximum error"], [
    ["Mass flow rate", "0.60% MAPE", "0.00715 kg/s", "2.03% APE"],
    ["Electrical power", "0.76% MAPE", "0.298 kW", "1.96% APE"],
    ["Discharge temperature", "0.88 °C MAE", "1.13 °C", "2.68 °C AE"],
], widths=[1.7, 1.4, 1.5, 1.5])
add_fig("parity_mass_flow.png", "Figure 4. Mass-flow parity on independent validation points.", width=5.0)
add_fig("parity_power.png", "Figure 5. Power-input parity on independent validation points.", width=5.0)
add_fig("parity_discharge_temperature.png", "Figure 6. Discharge-temperature parity on independent validation points.", width=5.0)
p("The parity points lie close to the perfect-prediction line and do not show strong proportional bias.")

p("9. Discussion", style="Heading 1")
p("The main strength of the model is separation of thermodynamics from regression: CoolProp determines state properties, while the polynomial represents efficiency behaviour. This is easier to audit than directly fitting mass flow and power as black-box outputs.")
p("Pressure ratio is the dominant variable. The decline in volumetric efficiency is consistent with clearance re-expansion in a reciprocating compressor. Isentropic efficiency falls as valve pressure drop, leakage, heat-transfer irreversibility and motor/inverter losses increase. Frequency has a smaller effect on efficiency but directly scales mass flow through displacement rate.")
p("The principal limitation is data provenance. The performance table is representative rather than a measured map for a named commercial compressor, so the low validation errors demonstrate internal consistency on smooth data rather than expected field accuracy. Genuine manufacturer data may show rating-standard differences, rounding and regions where a quadratic surface is inadequate. The model should not be extrapolated materially beyond β = 2.14–9.10 or 30–75 Hz without additional data.")
p("Recommended improvements are to import genuine manufacturer selection-software data, use measured suction/discharge pressures, test physics-guided clearance-volume models, add prediction intervals, and couple the compressor map to evaporator, condenser and expansion-device models.")

p("10. Conclusions", style="Heading 1")
for item in [
    "A complete MATLAB package was developed for a variable-speed hermetic reciprocating compressor using R134a.",
    "The compressor was sized for the supplied duty: displacement 1,147.6 cm³/rev, envelope pressure ratios 2.14 to 9.10, and the fitted model predicts 50.02 kW cooling at the design point (−10 °C / 45 °C, 5 K superheat, 50 Hz) with 20.33 kW input (cooling COP 2.46).",
    "Calculated volumetric efficiencies range from 0.462 to 0.865; isentropic efficiencies range from 0.530 to 0.797.",
    "The final two-dimensional models achieve R² = 0.9974 for ηv and R² = 0.9837 for ηis on fitting data.",
    "On 30 independent validation points, mass-flow MAPE is 0.60%, power MAPE is 0.76%, and discharge-temperature MAE is 0.88 °C.",
    "The MATLAB package is data-agnostic: a genuine manufacturer export can replace the representative CSV without code changes.",
]:
    doc.add_paragraph(item, style="List Number")

p("References", style="Heading 1")
refs = [
    "Wang, J. and Lu, W. “A Streamlined Polynomial Regression-Based Modeling of Speed-Driven Hermetic-Reciprocating Compressors.” Applied Sciences, 2025, 15(22), 12016.",
    "Bell, I. H., Wronski, J., Quoilin, S., and Lemort, V. “Pure and Pseudo-pure Fluid Thermophysical Property Evaluation and the Open-Source Thermophysical Property Library CoolProp.” Industrial & Engineering Chemistry Research, 2014, 53(6), 2498–2508.",
    "ASHRAE. ASHRAE Handbook—Refrigeration. American Society of Heating, Refrigerating and Air-Conditioning Engineers.",
    "Secop. BD50F Direct Current Compressor: R134a/R1234yf Technical Literature. Public manufacturer literature used only as background for variable-speed hermetic compressor practice.",
    "CoolProp documentation. Thermophysical property library and fluid property equations. http://www.coolprop.org/",
]
for ref in refs:
    doc.add_paragraph(ref, style="List Number")

p("Appendix A — MATLAB package outputs", style="Heading 1")
p("Running main_compressor_model.m produces calculated_variables.csv, regression_coefficients.csv, candidate_model_comparison.csv, one_dimensional_fits.csv, new_case_predictions.csv, validation_results.csv and the efficiency/parity figures in the results/figures folder.")

p("Appendix B — Sample calculated operating points", style="Heading 1")
sample_ids = [1, 3, 28, 54, 55, 84]
sub = calc[calc.point_id.isin(sample_ids)]
rows = []
for _, r in sub.iterrows():
    rows.append([int(r.point_id), r.dataset, f"{r.T_evap_C:g}", f"{r.T_cond_C:g}", f"{r.freq_Hz:g}", f"{r.beta:.2f}", f"{r.eta_v_calc:.3f}", f"{r.eta_is_calc:.3f}", f"{r.m_dot_kg_s:.5f}", f"{r.W_el_kW:.3f}", f"{r.T_dis_C:.1f}"])
add_table(["Point", "Dataset", "Tevap", "Tcond", "f", "β", "ηv", "ηis", "ṁ kg/s", "W kW", "Tdis °C"], rows, widths=[0.5,0.85,0.55,0.55,0.45,0.55,0.55,0.55,0.75,0.6,0.65], font_size=7.5)
p("Sample values are rounded; the CSV files are the authoritative numerical records.", italic=True)

doc.save(OUT)
print(f"Wrote {OUT}")
